import streamlit as st
import pandas as pd
import openpyxl
import altair as alt
import re
import unicodedata
import os
import base64
import html

from supabase import create_client, Client

# ---------- Configuration ----------
st.set_page_config(page_title="Bienvenu ! Votre salaire près de vous", page_icon="💰", layout="centered")

FICHIER_EXCEL = "ETAT_DES_PRIMES_Aout_2026.xlsx"
LOGO_FICHIER = "logo_xgs.png"
FEUILLES_IGNOREES = {"prime manager"}  # comparées en minuscule, sans accent

# Jetons de design. Le graphique et le CSS lisent les memes valeurs, pour que
# la couleur de serie du graphe et celle de l'interface ne divergent jamais.
POLICE = 'system-ui, -apple-system, "Segoe UI", Roboto, sans-serif'
SURFACE = "#ffffff"
ENCRE = "#0b0b0b"
ENCRE_2 = "#52514e"
ENCRE_3 = "#898781"
FILET = "#e1e0d9"
FILET_FORT = "#c3c2b7"
SERIE = "#2a78d6"   # bleu, fente categorielle 1 (validee sur surface blanche)
HAUSSE = "#006300"
BAISSE = "#d03b3b"

# ---------- Supabase ----------
@st.cache_resource
def initialiser_supabase():
    """
    Client public : authentification normale.
    Client admin : création des comptes sans envoi d'e-mail.
    La clé secrète reste uniquement côté serveur (Secrets Streamlit).
    """
    try:
        url = st.secrets["SUPABASE_URL"]
        publishable_key = st.secrets["SUPABASE_KEY"]
        secret_key = st.secrets["SUPABASE_SECRET_KEY"]

        if not url or not publishable_key or not secret_key:
            return None, None

        client = create_client(url, publishable_key)
        admin_client = create_client(url, secret_key)

        return client, admin_client

    except Exception:
        return None, None


supabase, supabase_admin = initialiser_supabase()


@st.cache_data
def charger_logo_base64():
    if not os.path.exists(LOGO_FICHIER):
        return None
    with open(LOGO_FICHIER, "rb") as f:
        return base64.b64encode(f.read()).decode()


def appliquer_style():
    """Le fond degrade et le logo restent la peau de l'outil ; la carte de contenu
    devient blanche opaque, pour que la surface sous le graphique soit connue
    (la couleur de serie est validee contre du blanc) et le texte bien contraste."""
    logo_b64 = charger_logo_base64()
    fond_logo = f', url("data:image/png;base64,{logo_b64}")' if logo_b64 else ""
    taille_logo = ", 42%" if logo_b64 else ""

    st.markdown(f"""
    <style>
    .stApp {{
        background-image: linear-gradient(135deg, rgba(10, 25, 90, 0.93) 0%, rgba(30, 70, 190, 0.90) 45%, rgba(70, 140, 240, 0.88) 100%){fond_logo};
        background-size: cover{taille_logo};
        background-position: center, center 18%;
        background-repeat: no-repeat, no-repeat;
        background-attachment: fixed, fixed;
    }}
    .block-container {{
        background-color: {SURFACE};
        border-radius: 18px;
        padding: 2.4rem 2.6rem 2rem;
        margin-top: 1.4rem;
        box-shadow: 0 6px 28px rgba(10, 25, 90, 0.28);
        max-width: 60rem;
    }}
    h1, h2, h3, h4, h5, p, label, li,
    .stMarkdown, .stCaption, .stTextInput label, .stSelectbox label {{
        color: {ENCRE};
    }}
    .block-container h1 {{
        font-size: 1.9rem;
        line-height: 1.2;
        letter-spacing: -0.01em;
        padding-top: 0;
    }}

    /* --- en-tete de l'agent --- */
    .entete-nom {{
        font-size: 1.35rem;
        font-weight: 600;
        color: {ENCRE};
    }}
    .puces {{ margin: .45rem 0 .2rem; }}
    .puce {{
        display: inline-block;
        border: 1px solid {FILET};
        border-radius: 999px;
        padding: .16rem .62rem;
        margin: 0 .3rem .3rem 0;
        font-size: .78rem;
        color: {ENCRE_2};
        white-space: nowrap;
    }}

    /* --- chiffre principal : un seul par ecran --- */
    .hero {{ margin: 0 0 1.3rem; }}
    .hero-label {{
        font-size: .85rem;
        color: {ENCRE_2};
        margin-bottom: .1rem;
    }}
    .hero-chiffre {{
        font-family: {POLICE};
        font-size: clamp(2.1rem, 7.5vw, 3.4rem);
        font-weight: 650;
        line-height: 1.05;
        color: {ENCRE};
    }}
    .hero-delta {{ font-size: .88rem; font-weight: 600; margin-top: .25rem; }}
    .hero-delta.hausse {{ color: {HAUSSE}; }}
    .hero-delta.baisse {{ color: {BAISSE}; }}
    .hero-delta.neutre {{ color: {ENCRE_3}; font-weight: 500; }}

    /* --- tuiles de detail --- */
    .tuile {{
        border: 1px solid {FILET};
        border-radius: 12px;
        padding: .8rem .95rem;
        height: 100%;
    }}
    .tuile-label {{
        font-size: .76rem;
        color: {ENCRE_3};
        line-height: 1.3;
    }}
    @media (min-width: 640px) {{
        .tuile-label {{ min-height: 2.1em; }}
    }}
    .tuile-valeur {{
        font-size: 1.3rem;
        font-weight: 600;
        color: {ENCRE};
        margin-top: .2rem;
    }}

    div[data-testid="stColumn"]:has(> div[data-testid="stVerticalBlock"]:empty) {{
        display: none;
    }}

    .titre-section {{
        font-size: 1.05rem;
        font-weight: 600;
        color: {ENCRE};
        margin: 1.6rem 0 .1rem;
    }}
    .note-section {{
        font-size: .82rem;
        color: {ENCRE_2};
        margin-bottom: .5rem;
    }}
    .pied {{ font-size: .78rem; color: {ENCRE_3}; }}
    </style>
    """, unsafe_allow_html=True)


appliquer_style()

MOIS_FR = [
    "janvier", "fevrier", "mars", "avril", "mai", "juin",

    "juillet", "aout", "septembre", "octobre", "novembre", "decembre"
]


def sans_accent(texte):
    if not isinstance(texte, str):
        return ""
    nfkd = unicodedata.normalize("NFD", texte)
    return "".join(c for c in nfkd if unicodedata.category(c) != "Mn").lower().strip()


def normaliser_matricule(valeur):
    """Forme de comparaison d'un matricule : sans espaces autour, en minuscules."""
    if valeur is None:
        return ""
    return str(valeur).strip().lower()


def extraire_mois_annee(nom_feuille):
    """Ex: 'Août 2026' -> ('Aout', 2026, index_mois)."""
    m = re.search(r"(\D+)\s*(\d{4})", nom_feuille)
    if not m:
        return nom_feuille, None, 99
    mot, annee = m.group(1).strip(), int(m.group(2))
    mot_norm = sans_accent(mot)
    idx = MOIS_FR.index(mot_norm) if mot_norm in MOIS_FR else 99
    return mot.strip().capitalize(), annee, idx


def trouver_ligne_entete(ws, max_lignes_scan=15):
    """Cherche la ligne contenant 'Matricule' ou 'Nom et prénoms'."""
    for r in range(1, max_lignes_scan + 1):
        valeurs = [sans_accent(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)]
        if "matricule" in valeurs or any("nom et prenom" in v for v in valeurs):
            return r
    return None


def mapper_colonnes(ws, ligne_entete, mois_norm):
    """Associe chaque champ standard à un numéro de colonne, via correspondance souple sur l'intitulé."""
    mapping = {}
    for c in range(1, ws.max_column + 1):
        intitule = sans_accent(ws.cell(row=ligne_entete, column=c).value)
        if not intitule:
            continue
        if "matricule" in intitule:
            mapping["Matricule"] = c
        elif "mot de passe" in intitule or "password" in intitule:
            mapping["MotDePasse"] = c
        elif "nom et prenom" in intitule:
            mapping["Nom"] = c
        elif intitule == "projet" or "projet" in intitule:
            mapping.setdefault("Projet", c)
        elif "superviseur" in intitule:
            mapping["Superviseur"] = c
        elif intitule == "poste":
            mapping["Poste"] = c
        elif "salaire de base" in intitule:
            mapping["Salaire_Base"] = c
        elif "net a payer" in intitule and "total" not in intitule:
            mapping["Net_A_Payer"] = c
        elif "coach" in intitule:
            mapping["Prime_Coach"] = c
        elif "prime" in intitule or "bonus" in intitule:
            # Le report du mois precedent n'a pas le meme intitule d'une feuille
            # a l'autre : "Bonus avril", "Prime Mai", "Prime Juillet"...
            mapping.setdefault("Prime_Anterieure", c)
        elif "total" in intitule:
            # "Total Juillet sans formule" doit etre teste avant "Total <mois>",
            # sinon il ecrase la colonne du total du mois sur la feuille Juillet.
            if "sans formule" in intitule:
                mapping.setdefault("Total_Mois_Precedent", c)
            elif mois_norm and mois_norm in intitule:
                mapping["Total_Mois"] = c
    return mapping


@st.cache_data(ttl=60)
def charger_donnees():
    if not os.path.exists(FICHIER_EXCEL):
        return None, []

    wb_valeurs = openpyxl.load_workbook(FICHIER_EXCEL, data_only=True)
    lignes = []
    feuilles_lues = []

    for nom_feuille in wb_valeurs.sheetnames:
        if sans_accent(nom_feuille) in FEUILLES_IGNOREES:
            continue

        mois_label, annee, ordre_mois = extraire_mois_annee(nom_feuille)
        if annee is None:
            continue  # feuille qui ne ressemble pas à un mois -> ignorée

        ws = wb_valeurs[nom_feuille]
        ligne_entete = trouver_ligne_entete(ws)
        if ligne_entete is None:
            continue

        mapping = mapper_colonnes(ws, ligne_entete, sans_accent(mois_label))
        if "Nom" not in mapping:
            continue

        feuilles_lues.append(nom_feuille)

        # L'intitule exact de la colonne de prime est conserve tel qu'il est
        # ecrit dans le classeur ("Bonus avril", "Prime Juillet"...) : c'est ce
        # libelle que l'agent reconnait, et il ne veut pas dire la meme chose
        # d'un mois a l'autre.
        colonne_prime = mapping.get("Prime_Anterieure")
        intitule_prime = ws.cell(row=ligne_entete, column=colonne_prime).value if colonne_prime else None
        intitule_prime = str(intitule_prime).strip() if intitule_prime not in (None, "") else None

        for r in range(ligne_entete + 1, ws.max_row + 1):
            nom = ws.cell(row=r, column=mapping["Nom"]).value
            if not nom or not str(nom).strip():
                continue  # ligne vide ou ligne de total en bas de tableau

            def val(champ):
                col = mapping.get(champ)
                return ws.cell(row=r, column=col).value if col else None

            matricule = val("Matricule")
            lignes.append({
                "Matricule": str(matricule).strip() if matricule not in (None, "") else None,
                "Nom": str(nom).strip(),
                "Mois": mois_label,
                "Annee": annee,
                "ordre_mois": ordre_mois,
                "Projet": val("Projet"),
                "Poste": val("Poste"),
                "Salaire_Base": val("Salaire_Base"),
                "Net_A_Payer": val("Net_A_Payer"),
                "Prime_Anterieure": val("Prime_Anterieure"),
                "Intitule_Prime": intitule_prime,
                "Prime_Coach": val("Prime_Coach"),
                "Total_Mois": val("Total_Mois"),
            })

    if not lignes:
        return None, feuilles_lues

    df = pd.DataFrame(lignes)
    df["Matricule_norm"] = df["Matricule"].map(normaliser_matricule)
    return df, feuilles_lues


def email_auth_depuis_matricule(matricule):
    """Identifiant technique utilisé par Supabase Auth à partir du matricule."""
    matricule_normalise = str(matricule).strip().lower()
    return f"{matricule_normalise}@xgs.local"


def creer_compte_supabase(matricule, mot_de_passe, nom=None):
    """
    Crée un compte Supabase avec l'API Admin.
    Aucun e-mail n'est envoyé.
    email_confirm=True considère immédiatement l'identifiant comme confirmé.
    """
    email = email_auth_depuis_matricule(matricule)

    return supabase_admin.auth.admin.create_user({
        "email": email,
        "password": str(mot_de_passe),
        "email_confirm": True,
        "user_metadata": {
            "matricule": str(matricule).strip(),
            "nom": str(nom or "").strip(),
        },
    })


def authentifier_ou_creer_compte(matricule, mot_de_passe, nom=None):
    """
    Authentifie l'utilisateur.
    
    - Si le compte existe : connexion normale.
    - Si le compte n'existe pas : création du compte sans email,
      puis connexion immédiate.
    """

    if supabase is None or supabase_admin is None:
        return False, (
            "Supabase n'est pas configuré. Vérifiez les secrets de "
            "l'application."
        )

    matricule = str(matricule).strip()
    mot_de_passe = str(mot_de_passe)

    if not matricule:
        return False, "Veuillez saisir votre matricule."

    if not mot_de_passe:
        return False, "Veuillez saisir votre mot de passe."

    email = email_auth_depuis_matricule(matricule)

    # ============================================================
    # 1. VÉRIFIER SI LE COMPTE EXISTE DÉJÀ
    # ============================================================

    utilisateur_existant = None

    try:
        utilisateurs = supabase_admin.auth.admin.list_users()

        # Compatibilité avec différentes versions de supabase-py
        if hasattr(utilisateurs, "users"):
            utilisateurs = utilisateurs.users

        elif hasattr(utilisateurs, "data"):
            if hasattr(utilisateurs.data, "users"):
                utilisateurs = utilisateurs.data.users
            elif isinstance(utilisateurs.data, list):
                utilisateurs = utilisateurs.data

        for utilisateur in utilisateurs or []:

            if getattr(utilisateur, "email", "").lower() == email.lower():
                utilisateur_existant = utilisateur
                break

    except Exception as e:

        return False, f"Impossible de vérifier le compte utilisateur : {e}"

    # ============================================================
    # 2. LE COMPTE EXISTE → CONNEXION
    # ============================================================

    if utilisateur_existant is not None:

        try:

            supabase.auth.sign_in_with_password({
                "email": email,
                "password": mot_de_passe
            })

            return True, "Connexion réussie."

        except Exception:

            return False, "Mot de passe incorrect."

    # ============================================================
    # 3. LE COMPTE N'EXISTE PAS → PREMIÈRE CONNEXION
    # ============================================================

    try:

        supabase_admin.auth.admin.create_user({
            "email": email,
            "password": mot_de_passe,
            "email_confirm": True,
            "user_metadata": {
                "matricule": matricule,
                "nom": str(nom or "").strip()
            }
        })

    except Exception as e:

        return False, f"Impossible de créer le compte : {e}"

    # ============================================================
    # 4. CONNEXION IMMÉDIATE APRÈS CRÉATION
    # ============================================================

    try:

        supabase.auth.sign_in_with_password({
            "email": email,
            "password": mot_de_passe
        })

        return True, "Compte créé et connexion réussie."

    except Exception as e:

        return False, f"Compte créé mais connexion impossible : {e}"

def formater_fcfa(valeur):
    """Montant en FCFA. Une cellule vide du classeur arrive ici en NaN, qui se
    formate sans lever d'exception : il faut donc le tester avant."""
    if pd.isna(valeur) or valeur == "":
        return "-"
    try:
        return f"{valeur:,.0f} FCFA".replace(",", " ")
    except (ValueError, TypeError):
        return valeur

# ---------- Briques d'affichage ----------
def libelle_prime(ligne):
    """L'intitule du classeur ("Bonus avril", "Prime Juillet") plutot qu'un
    libelle fige : il ne designe pas le meme mois d'une feuille a l'autre."""
    intitule = ligne.get("Intitule_Prime")
    if intitule in (None, "") or pd.isna(intitule):
        return "Prime du mois precedent"
    return str(intitule)


def tuile(label, valeur):
    return (
        '<div class="tuile">'
        f'<div class="tuile-label">{html.escape(str(label))}</div>'
        f'<div class="tuile-valeur">{html.escape(str(valeur))}</div>'
        "</div>"
    )


def bloc_ecart(total, total_precedent, mois_precedent):
    """Ecart avec le mois precedent. La fleche et le signe portent le sens ;
    la couleur ne fait que le confirmer."""
    if pd.isna(total):
        return '<div class="hero-delta neutre">Total du mois non renseign&eacute; dans le fichier</div>'

    if total_precedent is None:
        return '<div class="hero-delta neutre">Premier mois pay&eacute; de l&rsquo;ann&eacute;e</div>'

    if pd.isna(total_precedent):
        return (
            '<div class="hero-delta neutre">Pas de comparaison : le total du mois '
            f"précédent ({html.escape(mois_precedent)}) n&rsquo;est pas renseigné</div>"
        )

    ecart = float(total) - float(total_precedent)
    if abs(ecart) < 1:
        return f'<div class="hero-delta neutre">Identique à {html.escape(mois_precedent)}</div>'

    fleche, classe, signe = ("▲", "hausse", "+") if ecart > 0 else ("▼", "baisse", "−")
    return (
        f'<div class="hero-delta {classe}">{fleche} {signe}{formater_fcfa(abs(ecart))} '
        f"par rapport à {html.escape(mois_precedent)}</div>"
    )


def graphique_evolution(vue_annee, mois_selectionne):
    """Evolution du total paye, mois par mois. Une seule serie : pas de legende,
    le titre de section dit ce qui est trace. Le tableau en dessous sert de
    version lisible sans couleur."""
    donnees = vue_annee.dropna(subset=["Total_Mois"]).copy()
    if len(donnees) < 2:
        return None

    donnees["Montant"] = donnees["Total_Mois"].astype(float)
    donnees["Montant_FCFA"] = donnees["Total_Mois"].apply(formater_fcfa)
    ordre = donnees.sort_values("ordre_mois")["Mois"].tolist()

    axe_x = alt.X(
        "Mois:N",
        sort=ordre,
        title=None,
        axis=alt.Axis(
            labelAngle=0, labelColor=ENCRE_2, labelFontSize=12, labelPadding=8,
            domainColor=FILET_FORT, tickColor=FILET_FORT, grid=False,
        ),
    )
    axe_y = alt.Y(
        "Montant:Q",
        title=None,
        scale=alt.Scale(zero=True, nice=True),
        axis=alt.Axis(
            labelColor=ENCRE_3, labelFontSize=11, labelPadding=8,
            gridColor=FILET, gridWidth=1, domain=False, ticks=False,
            labelExpr="replace(format(datum.value, ',.0f'), /,/g, ' ')",
        ),
    )
    infobulle = [
        alt.Tooltip("Mois:N", title="Mois"),
        alt.Tooltip("Montant_FCFA:N", title="Total du mois"),
    ]

    base = alt.Chart(donnees)
    repere = (
        alt.Chart(pd.DataFrame({"Mois": [mois_selectionne]}))
        .mark_rule(color=FILET_FORT, strokeWidth=1)
        .encode(x=axe_x)
    )
    aire = base.mark_area(color=SERIE, opacity=0.10, line=False).encode(x=axe_x, y=axe_y)
    ligne_serie = base.mark_line(
        color=SERIE, strokeWidth=2, strokeJoin="round", strokeCap="round"
    ).encode(x=axe_x, y=axe_y)
    points = base.mark_point(
        color=SERIE, filled=True, size=90, stroke=SURFACE, strokeWidth=2
    ).encode(x=axe_x, y=axe_y, tooltip=infobulle)
    # Une seule etiquette directe, sur le mois selectionne : elle porte donc le
    # meme montant que le chiffre principal. L'infobulle et le tableau portent
    # les autres valeurs.
    etiquette = (
        alt.Chart(donnees[donnees["Mois"] == mois_selectionne])
        .mark_text(dy=-18, align="center", fontSize=13, fontWeight=600, color=ENCRE)
        .encode(x=axe_x, y=axe_y, text="Montant_FCFA:N")
    )

    return (
        (repere + aire + ligne_serie + points + etiquette)
        .properties(height=260, padding={"top": 18, "right": 16, "bottom": 4, "left": 8})
        .configure_view(strokeWidth=0)
        .configure(background="transparent", font=POLICE)
        .configure_axis(labelFont=POLICE, titleFont=POLICE)
        .configure_axisX(grid=False)
        .configure_text(font=POLICE)
    )


# ---------- Interface ----------
df, feuilles_lues = charger_donnees()

if df is None:
    st.error(
        f"Impossible de trouver des données exploitables dans '{FICHIER_EXCEL}'. "
        "Vérifiez que le fichier est présent et que la colonne Matricule est bien remplie."
    )
    st.stop()


# ---------- Ecran de connexion ----------
if not st.session_state.get("connecte"):
    st.title("Bienvenu ! Votre salaire près de vous")
    st.caption(
        "Entrez votre matricule et votre mot de passe pour consulter le détail, mois par mois. "
        "Lors de votre première connexion, votre compte est créé automatiquement."
    )

    with st.form("connexion"):
        matricule_saisi = st.text_input("Votre matricule", placeholder="Ex : XGS-TC01").strip()
        mot_de_passe_saisi = st.text_input("Votre mot de passe", type="password").strip()
        valider = st.form_submit_button("Voir mon salaire")

    if valider:
        if not matricule_saisi or not mot_de_passe_saisi:
            st.warning("Veuillez renseigner votre matricule et votre mot de passe.")
        else:
            correspondances = df[
                df["Matricule_norm"] == normaliser_matricule(matricule_saisi)
            ]

            if correspondances.empty:
                st.error("Matricule inconnu. Contactez le service RH si besoin.")
            else:
                valeur_nom = correspondances.iloc[0]["Nom"]
                nom_utilisateur = "" if pd.isna(valeur_nom) else str(valeur_nom).strip()

                authentifie, message = authentifier_ou_creer_compte(
                    matricule_saisi,
                    mot_de_passe_saisi,
                    nom_utilisateur
                )

                if authentifie:
                    st.session_state["connecte"] = True
                    st.session_state["matricule_connecte"] = normaliser_matricule(matricule_saisi)
                    st.session_state["matricule_affiche"] = str(
                        correspondances.iloc[0]["Matricule"]
                    ).strip()
                    st.session_state["message_connexion"] = message
                    st.rerun()
                else:
                    st.error(message)

    st.divider()
    st.markdown(
        '<div class="pied">Cette page est en lecture seule : les données de salaire sont '
        "consultées depuis le fichier de référence. L’authentification est gérée par "
        "Supabase.</div>",
        unsafe_allow_html=True,
    )
    st.stop()


# ---------- Ecran de l'agent ----------
resultats = df[df["Matricule_norm"] == st.session_state["matricule_connecte"]].copy()

if resultats.empty:
    st.warning(
        "Aucune donnée trouvée pour ce matricule. "
        "Vérifiez la saisie, ou contactez le service RH si votre matricule n'a pas encore été renseigné."
    )
    if st.button("Se déconnecter"):
        st.session_state.clear()
        st.rerun()
    st.stop()

# Un matricule partage par deux personnes ferait voir a chacune le salaire de
# l'autre, et rien dans la saisie ne permet de les distinguer. On refuse donc
# d'afficher quoi que ce soit plutot que de melanger deux paies.
noms_distincts = sorted(
    {str(n).strip() for n in resultats["Nom"].dropna() if str(n).strip()}
)
if len(noms_distincts) > 1:
    st.error(
        "Ce matricule est associé à plusieurs noms dans le fichier de référence : "
        + ", ".join(noms_distincts)
        + ". Par précaution, aucun salaire n'est affiché. Contactez le service RH "
        "pour faire corriger le matricule."
    )
    if st.button("Se déconnecter"):
        st.session_state.clear()
        st.rerun()
    st.stop()

message_connexion = st.session_state.pop("message_connexion", None)
if message_connexion:
    st.success(message_connexion)

resultats = resultats.sort_values(["Annee", "ordre_mois"])
premiere_ligne = resultats.iloc[0]
nom_complet = "" if pd.isna(premiere_ligne["Nom"]) else str(premiere_ligne["Nom"]).strip()

puces = [f"Matricule {st.session_state.get('matricule_affiche', '')}"]
for champ in ("Projet", "Poste"):
    valeur = premiere_ligne.get(champ)
    if not pd.isna(valeur) and str(valeur).strip():
        puces.append(str(valeur).strip())

st.markdown(
    f'<div class="entete-nom">Bonjour, {html.escape(nom_complet)}</div>'
    '<div class="puces">'
    + "".join(f'<div class="puce">{html.escape(p)}</div>' for p in puces)
    + "</div>",
    unsafe_allow_html=True,
)

# --- une seule rangee de filtres, au-dessus de tout ce qu'elle cadre ---
annees_disponibles = sorted(resultats["Annee"].dropna().unique().tolist(), reverse=True)
colonnes_filtres = st.columns([1, 1, 2])

if len(annees_disponibles) > 1:
    annee_choisie = colonnes_filtres[0].selectbox("Année", annees_disponibles)
    emplacement_mois = colonnes_filtres[1]
else:
    annee_choisie = annees_disponibles[0]
    emplacement_mois = colonnes_filtres[0]

vue_annee = resultats[resultats["Annee"] == annee_choisie]
mois_disponibles = vue_annee["Mois"].tolist()
mois_choisi = emplacement_mois.selectbox(
    "Mois", mois_disponibles, index=len(mois_disponibles) - 1
)

rang = mois_disponibles.index(mois_choisi)
ligne = vue_annee.iloc[rang]
ligne_precedente = vue_annee.iloc[rang - 1] if rang > 0 else None

# --- le chiffre que l'agent vient chercher : un seul par ecran ---
st.markdown(
    '<div class="hero">'
    f'<div class="hero-label">Total payé — {html.escape(mois_choisi)} {annee_choisie}</div>'
    f'<div class="hero-chiffre">{formater_fcfa(ligne["Total_Mois"])}</div>'
    + bloc_ecart(
        ligne["Total_Mois"],
        None if ligne_precedente is None else ligne_precedente["Total_Mois"],
        "" if ligne_precedente is None else str(ligne_precedente["Mois"]),
    )
    + "</div>",
    unsafe_allow_html=True,
)

# --- de quoi ce total est fait ---
colonnes_tuiles = st.columns(4)
detail_du_mois = [
    ("Salaire de base", ligne["Salaire_Base"]),
    ("Net à payer (hors primes)", ligne["Net_A_Payer"]),
    (libelle_prime(ligne), ligne["Prime_Anterieure"]),
    ("Prime coach métier", ligne["Prime_Coach"]),
]
for colonne, (label, valeur) in zip(colonnes_tuiles, detail_du_mois):
    colonne.markdown(tuile(label, formater_fcfa(valeur)), unsafe_allow_html=True)

# --- evolution mois par mois ---
st.markdown('<div class="titre-section">Évolution du total payé</div>', unsafe_allow_html=True)
graphique = graphique_evolution(vue_annee, mois_choisi)

if graphique is None:
    st.markdown(
        '<div class="note-section">Un seul mois payé pour le moment : l’évolution '
        "s’affichera dès le mois suivant.</div>",
        unsafe_allow_html=True,
    )
else:
    st.markdown(
        f'<div class="note-section">Mois par mois sur {annee_choisie}. Le repère vertical '
        f"marque {html.escape(mois_choisi)}.</div>",
        unsafe_allow_html=True,
    )
    st.altair_chart(graphique, use_container_width=True, theme=None)

cumul = vue_annee["Total_Mois"].sum(min_count=1)
nombre_mois = int(vue_annee["Total_Mois"].notna().sum())
if nombre_mois:
    st.markdown(
        f'<div class="note-section">Cumul {annee_choisie} sur {nombre_mois} mois '
        f"payé{'s' if nombre_mois > 1 else ''} : "
        f"<strong>{formater_fcfa(cumul)}</strong></div>",
        unsafe_allow_html=True,
    )

# --- tableau complet : la version du graphique lisible sans couleur ---
st.markdown('<div class="titre-section">Tableau détaillé</div>', unsafe_allow_html=True)

colonnes_affichees = [
    "Mois", "Salaire_Base", "Net_A_Payer", "Prime_Anterieure", "Prime_Coach", "Total_Mois",
]
noms_lisibles = {
    "Mois": "Mois",
    "Salaire_Base": "Salaire de base",
    "Net_A_Payer": "Net à payer (hors primes)",
    "Prime_Anterieure": "Prime / bonus",
    "Prime_Coach": "Prime coach métier",
    "Total_Mois": "Total du mois",
}
tableau = vue_annee[colonnes_affichees].rename(columns=noms_lisibles).reset_index(drop=True)
for colonne in tableau.columns:
    if colonne != "Mois":
        tableau[colonne] = tableau[colonne].apply(formater_fcfa)

st.dataframe(tableau, hide_index=True, use_container_width=True)

st.divider()
if st.button("Se déconnecter"):
    st.session_state.clear()
    st.rerun()

st.markdown(
    '<div class="pied">Cette page est en lecture seule : les données de salaire sont '
    "consultées depuis le fichier de référence. L’authentification est gérée par "
    "Supabase.</div>",
    unsafe_allow_html=True,
)
